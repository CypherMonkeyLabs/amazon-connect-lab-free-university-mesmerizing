#!/usr/bin/env python3
"""Build the Amazon Connect round-2 static lab from the cited curriculum."""

from __future__ import annotations

import html
import json
import os
import re
import shutil
import signal
import subprocess
import sys
import tempfile
import threading
import time
import zipfile
from collections import defaultdict
from pathlib import Path

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "content"))
from curriculum import LESSONS, TRACKS  # noqa: E402
from glossary import GLOSSARY  # noqa: E402
from iac import IAC  # noqa: E402

SITE = "https://cyphermonkeylabs.com"
CONTACT = SITE + "/#contact"
ORG = "https://github.com/CypherMonkeyLabs"
REPO = ORG + "/amazon-connect-lab-free-university-mesmerizing"
TREE = "https://linktr.ee/cypherMonkey"
DISCLAIMER = (
    "This is an independent, free, learning-only student project. It respects every company and vendor. "
    "Trademarks belong to their owners. It is not affiliated with or endorsed by Amazon Web Services, Cisco, "
    "Genesys, or any other company. If you have a concern, the page will be removed or corrected promptly "
    "through the contact form. Nothing here is sold. There are no ads. The author is happy to volunteer for any organization."
)
COPY = "© 2026 CypherMonkeyLabs 🐒"
FILLS = ["#e0e7ff", "#cffafe", "#fce7f3", "#fef3c7", "#d1fae5", "#ffedd5"]
STROKES = ["#4338ca", "#0e7490", "#9d174d", "#b45309", "#047857", "#c2410c"]

BANNER = [
    ("home", "Home", "index.html", "🏠", "Front door and the track map"),
    ("learn", "Learn", "start.html", "📘", "Ten-minute path and the full course"),
    ("diagrams", "Diagrams", "diagrams.html", "🗺️", "Every picture as SVG, PNG, and HTML"),
    ("flows", "Flows", "flows.html", "〰️", "Animated call and data paths"),
    ("play", "Play", "play.html", "🎮", "Routing, Lambda, port, and bridge drills"),
    ("listen", "Listen", "listen.html", "🎧", "Study clips and read-along"),
    ("downloads", "Downloads", "downloads.html", "⬇️", "Sheets, cards, templates, and the zip"),
    ("guide", "Guide", "help.html", "🐒", "Menu, course, bot, and downloads"),
]
DRAWER = [
    ("Study", [
        ("home", "Home", "index.html", "🏠", "Front door and the track map"),
        ("learn", "Learn", "start.html", "📘", "Ten-minute path and the full course"),
        ("rooms", "All rooms", "index.html#rooms", "🗂️", "Eighteen tracks, one room at a time"),
        ("mastery", "Mastery", "mastery.html", "🧭", "Fundamentals through the bridge"),
        ("ecosystem", "Ecosystem", "ecosystem.html", "🌐", "Glossary, map, bridges, and hooks"),
    ]),
    ("Pictures", [
        ("diagrams", "Diagrams", "diagrams.html", "🗺️", "Every picture as SVG, PNG, and HTML"),
        ("flows", "Flows", "flows.html", "〰️", "Animated call and data paths"),
        ("charts", "Charts", "charts.html", "📊", "Counts from this repository"),
    ]),
    ("Practice", [
        ("play", "Play", "play.html", "🎮", "Routing, Lambda, port, and bridge drills"),
        ("quizzes", "Quizzes", "quizzes.html", "✅", "Practice questions with cited answers"),
        ("cards", "Flashcards", "flashcards.html", "🃏", "Know or again, stored in this browser"),
        ("interviews", "Interviews", "interviews.html", "🎤", "Fictional practice. Not a biography"),
        ("listen", "Listen", "listen.html", "🎧", "Study clips and read-along"),
    ]),
    ("Reference", [
        ("downloads", "Downloads", "downloads.html", "⬇️", "Sheets, cards, templates, and the zip"),
        ("glossary", "Glossary", "glossary.html", "📖", "Terms from the written rooms"),
        ("ports", "Ports", "ports.html", "🔌", "TCP 443 and UDP 3478 from the network lesson"),
        ("regions", "Regions", "regions.html", "🗺️", "The published Global Resiliency pairs"),
        ("certs", "Certifications", "certs.html", "🎓", "What the catalog publishes, and what it does not"),
        ("careers", "Careers", "careers.html", "💼", "Study paths. Not a job offer"),
    ]),
    ("About", [
        ("guide", "Help", "help.html", "🐒", "Menu, course, bot, and downloads"),
        ("legal", "Legal", "legal.html", "🛡️", "Copyright, privacy, and removal"),
    ]),
]


def esc(value) -> str:
    return html.escape(str(value if value is not None else ""), quote=True)


def rel(depth: int, href: str) -> str:
    if href.startswith(("http://", "https://", "#")):
        return href
    return "../" * depth + href


def lesson_map():
    return {item["id"]: item for item in LESSONS}


def by_track():
    grouped = defaultdict(list)
    for item in LESSONS:
        grouped[item["track"]].append(item)
    return grouped


def confirm_needed(lesson) -> bool:
    blob = " ".join(lesson["paras"] + lesson["points"] + [lesson["why"], lesson["mistake"]]).lower()
    return any(word in blob for word in ("price", "usd", "quota", "badge", "exam", "rate"))


def load_quiz():
    path = ROOT / "campus" / "data" / "quiz.json"
    if not path.exists():
        return []
    return json.loads(path.read_text(encoding="utf-8"))["questions"]


def pick_quiz(lesson, pool):
    words = set(re.findall(r"[a-z]{4,}", (lesson["title"] + " " + lesson["why"]).lower()))
    ranked = sorted(
        pool,
        key=lambda item: -len(words & set(re.findall(r"[a-z]{4,}", item["q"].lower()))),
    )
    chosen = ranked[:3] if ranked else []
    if len(chosen) >= 3:
        return chosen
    return [{
        "q": "Which caution belongs to this room?",
        "choices": [lesson["mistake"], "Ignore the cited source", "Treat a lab estimate as an invoice", "Skip the error branch"],
        "answer": 0,
        "why": lesson["mistake"],
        "url": lesson["cite"]["url"],
        "source": lesson["cite"]["title"],
    }]


def wrap_lines(text, width=18):
    words = str(text).split()
    lines, cur = [], ""
    for word in words:
        trial = (cur + " " + word).strip()
        if len(trial) > width and cur:
            lines.append(cur)
            cur = word
        else:
            cur = trial
    if cur:
        lines.append(cur)
    return lines[:4] or [""]


def text_block(x, y, w, h, label, index):
    lines = wrap_lines(label)
    tspans = []
    start = y + h / 2 - (len(lines) - 1) * 8
    for i, line in enumerate(lines):
        tspans.append(f'<tspan x="{x + w / 2}" y="{start + i * 16}">{esc(line)}</tspan>')
    return (
        f'<rect x="{x}" y="{y}" width="{w}" height="{h}" rx="16" fill="{FILLS[index % 6]}" stroke="{STROKES[index % 6]}" stroke-width="3"/>'
        f'<text text-anchor="middle" font-family="system-ui,sans-serif" font-size="14" font-weight="700" fill="#1e1b4b">{"".join(tspans)}</text>'
    )


def svg_doc(width, height, title, body, caption):
    return (
        f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 {width} {height}" role="img" aria-label="{esc(title)}">'
        f"<title>{esc(title)}</title>"
        f'<rect width="100%" height="100%" fill="#fff7ed"/>'
        f"{body}"
        f'<text x="24" y="{height - 18}" font-family="system-ui,sans-serif" font-size="13" fill="#4338ca">{esc(caption)}</text>'
        "</svg>"
    )


def path_svg(title, nodes, caption, animated):
    count = max(1, len(nodes))
    width = 40 + count * 180
    height = 280
    parts = []
    motion = []
    for i, node in enumerate(nodes):
        x = 24 + i * 180
        parts.append(text_block(x, 70, 150, 84, node["label"], i))
        if i:
            x1 = 24 + (i - 1) * 180 + 150
            x2 = x
            parts.append(f'<path d="M {x1} 112 H {x2}" fill="none" stroke="#7c3aed" stroke-width="4"/>')
            motion.append(f"M {x1} 112 H {x2}")
    if animated and motion:
        parts.append(
            '<circle r="8" fill="#db2777"><animateMotion dur="4.5s" repeatCount="indefinite" path="'
            + " ".join(motion) + '"/></circle>'
        )
    return svg_doc(width, height, title, "".join(parts), caption)


def layers_svg(title, nodes, caption):
    height = 80 + len(nodes) * 78
    parts = []
    for i, node in enumerate(nodes):
        y = 24 + i * 78
        parts.append(text_block(80, y, 520, 64, node["label"], i))
    return svg_doc(680, height, title, "".join(parts), caption)


def split_svg(title, nodes, caption):
    left = nodes[0]["label"] if nodes else ""
    right = nodes[1]["label"] if len(nodes) > 1 else ""
    body = text_block(24, 70, 280, 110, left, 0) + text_block(360, 70, 280, 110, right, 2)
    body += '<path d="M 304 125 H 360" fill="none" stroke="#7c3aed" stroke-width="4"/>'
    return svg_doc(680, 280, title, body, caption)


def bars_svg(title, nodes, caption):
    values = []
    for node in nodes:
        match = re.search(r"[\d.]+", str(node.get("value", node["label"])))
        values.append(float(match.group()) if match else 1)
    peak = max(values) or 1
    parts = []
    for i, node in enumerate(nodes):
        y = 30 + i * 48
        width = 40 + (values[i] / peak) * 460
        parts.append(f'<rect x="180" y="{y}" width="{width:.1f}" height="32" rx="8" fill="{FILLS[i % 6]}" stroke="{STROKES[i % 6]}"/>')
        parts.append(f'<text x="16" y="{y + 22}" font-family="system-ui,sans-serif" font-size="14" fill="#1e1b4b">{esc(node["label"])}</text>')
    return svg_doc(760, 50 + len(nodes) * 48, title, "".join(parts), caption)


def diagram_svg(title, diagram, animated=False):
    kind = diagram.get("type")
    nodes = diagram.get("nodes") or []
    caption = diagram.get("caption") or ""
    if kind == "layers":
        return layers_svg(title, nodes, caption)
    if kind == "split":
        return split_svg(title, nodes, caption)
    if kind == "bars":
        return bars_svg(title, nodes, caption)
    return path_svg(title, nodes, caption, animated)


def flow_svg(title, steps):
    count = max(1, len(steps))
    height = 70 + count * 110
    parts = []
    centers = []
    for i, step in enumerate(steps):
        y = 24 + i * 110
        parts.append(text_block(180, y, 400, 78, step["label"], i))
        parts.append(
            f'<text x="180" y="{y + 96}" font-family="system-ui,sans-serif" font-size="13" fill="#4338ca">{esc(step["detail"])}</text>'
        )
        centers.append((380, y + 39))
        if i:
            parts.append(f'<path d="M 380 {y - 32} V {y}" fill="none" stroke="#7c3aed" stroke-width="4"/>')
    if len(centers) > 1:
        path = f"M {centers[0][0]} {centers[0][1]} " + " ".join(f"L {x} {y}" for x, y in centers[1:])
        parts.append(f'<circle r="8" fill="#db2777"><animateMotion dur="5s" repeatCount="indefinite" path="{path}"/></circle>')
    return svg_doc(760, height, title, "".join(parts), "Animated study flow. Confirm the cited source before relying on it.")


def mind_svg(track, lessons):
    title = track["title"] + " mind map"
    body = ['<circle cx="360" cy="210" r="70" fill="#ede9fe" stroke="#7c3aed" stroke-width="3"/>',
            f'<text x="360" y="214" text-anchor="middle" font-family="system-ui,sans-serif" font-size="16" font-weight="800" fill="#1e1b4b">{esc(track["title"])}</text>']
    spots = [(80, 60), (640, 60), (70, 340), (650, 340), (360, 40), (360, 380)]
    for i, lesson in enumerate(lessons[:6]):
        x, y = spots[i]
        body.append(text_block(x, y, 160, 64, lesson["title"], i))
    return svg_doc(860, 460, title, "".join(body), track["why"])


SHIELD = '<svg viewBox="0 0 48 48" aria-hidden="true"><path fill="currentColor" d="M24 4 8 10v12c0 10 6.8 16.8 16 20 9.2-3.2 16-10 16-20V10z"/><path fill="#fff" d="m16 24 5 5 11-12" stroke="#fff" stroke-width="2"/></svg>'
MONKEY_ICON = '<svg viewBox="0 0 64 64" aria-hidden="true"><circle cx="18" cy="24" r="8" fill="#fbbf24"/><circle cx="46" cy="24" r="8" fill="#fbbf24"/><circle cx="32" cy="36" r="16" fill="#fde68a"/><rect x="20" y="30" width="10" height="8" rx="2" fill="none" stroke="#1e1b4b" stroke-width="2"/><rect x="34" y="30" width="10" height="8" rx="2" fill="none" stroke="#1e1b4b" stroke-width="2"/><path d="M30 34h4" stroke="#1e1b4b" stroke-width="2"/><circle cx="25" cy="34" r="1.3" fill="#1e1b4b"/><circle cx="39" cy="34" r="1.3" fill="#1e1b4b"/><path d="M26 42q6 4 12 0" fill="none" stroke="#1e1b4b" stroke-width="1.8" stroke-linecap="round"/></svg>'
GITHUB_ICON = '<svg viewBox="0 0 24 24" aria-hidden="true"><path fill="currentColor" d="M12 .5A11.5 11.5 0 0 0 .5 12.3c0 5.2 3.4 9.6 8.1 11.2.6.1.8-.3.8-.6v-2c-3.3.7-4-1.4-4-1.4-.5-1.4-1.3-1.7-1.3-1.7-1-.7.1-.7.1-.7 1.2.1 1.8 1.2 1.8 1.2 1 1.8 2.8 1.3 3.4 1 .1-.8.4-1.3.8-1.6-2.6-.3-5.4-1.3-5.4-6 0-1.3.5-2.4 1.2-3.2-.1-.3-.5-1.5.1-3.2 0 0 1-.3 3.3 1.2a11 11 0 0 1 6 0c2.3-1.5 3.3-1.2 3.3-1.2.6 1.7.2 2.9.1 3.2.8.8 1.2 1.9 1.2 3.2 0 4.7-2.8 5.7-5.4 6 .4.4.8 1.1.8 2.3v3.3c0 .3.2.7.8.6A11.5 11.5 0 0 0 23.5 12.3 11.5 11.5 0 0 0 12 .5z"/></svg>'
TREE_ICON = '<svg viewBox="0 0 64 64" aria-hidden="true"><circle cx="32" cy="11" r="7" fill="#052e16"/><path d="M32 18 V27" stroke="#052e16" stroke-width="5" stroke-linecap="round"/><path d="M13 36 C22 25 42 25 51 36" fill="none" stroke="#052e16" stroke-width="5" stroke-linecap="round"/><circle cx="13" cy="38" r="7" fill="#052e16"/><circle cx="51" cy="38" r="7" fill="#052e16"/><circle cx="32" cy="28" r="5.5" fill="#052e16"/><path d="M32 33 V49" stroke="#052e16" stroke-width="6" stroke-linecap="round"/><path d="M19 55 H45" stroke="#052e16" stroke-width="6" stroke-linecap="round"/></svg>'


def footer_block():
    return f'''<footer class="site-footer" data-canon-footer>
<div class="disclaimer-panel" role="note">
<div class="disclaimer-ico" aria-hidden="true">{SHIELD}</div>
<div>
<h2>Independent study lab</h2>
<p>{esc(DISCLAIMER)}</p>
<p>Contact: <a href="{CONTACT}">{CONTACT}</a></p>
</div>
</div>
<p class="footer-brand">CypherMonkey Labs</p>
<nav class="brand-row" aria-label="CypherMonkey Labs">
<a class="brand-icon brand-web" href="{SITE}" aria-label="CypherMonkey Labs website"><span class="sr">Website</span>{MONKEY_ICON}</a>
<a class="brand-icon brand-gh" href="{ORG}" aria-label="CypherMonkey Labs on GitHub"><span class="sr">GitHub</span>{GITHUB_ICON}</a>
<a class="brand-icon brand-tree" href="{TREE}" aria-label="Linktree"><span class="sr">Linktree</span>{TREE_ICON}</a>
</nav>
<p class="site-copy">{COPY}</p>
</footer>'''


def drawer_html(depth, current):
    parts = ['<div class="drawer-bar"><p>Menu</p><button type="button" class="drawer-close" data-close-drawer>Close</button></div>',
             '<label class="menu-search"><span class="sr">Search the menu</span><input type="search" data-menu-filter placeholder="Search the menu" aria-label="Search the menu"></label>']
    for group, items in DRAWER:
        parts.append(f'<p class="drawer-group">{esc(group)}</p>')
        for key, label, href, icon, desc in items:
            current_attr = ' aria-current="page"' if key == current else ""
            parts.append(
                f'<a class="drawer-row" href="{esc(rel(depth, href))}"{current_attr}>'
                f'<span class="drawer-ico" aria-hidden="true">{icon}</span>'
                f"<span><strong>{esc(label)}</strong><small>{esc(desc)}</small></span></a>"
            )
    return "".join(parts)


def banner_html(depth, current):
    chips = []
    for key, label, href, icon, desc in BANNER:
        attr = ' aria-current="page"' if key == current else ""
        tour = f' data-tour="{key}"' if key in {"home", "learn", "diagrams", "play", "listen", "downloads", "guide"} else ""
        chips.append(
            f'<a class="section-chip" href="{esc(rel(depth, href))}" title="{esc(desc)}"{attr}{tour}>'
            f'<span class="chip-ico" aria-hidden="true">{icon}</span>{esc(label)}</a>'
        )
    return "".join(chips)


def dock_html(depth):
    items = [
        ("learn", "Learn", "start.html", "📘"),
        ("diagrams", "Diagrams", "diagrams.html", "🗺️"),
        ("play", "Play", "play.html", "🎮"),
        ("listen", "Listen", "listen.html", "🎧"),
        ("guide", "Guide", "help.html", "🐒"),
    ]
    parts = ['<button type="button" data-open-menu data-tour="menu"><span class="dock-ico" aria-hidden="true">☰</span>Menu</button>']
    for key, label, href, icon in items:
        parts.append(
            f'<a data-tour="{key}" href="{esc(rel(depth, href))}"><span class="dock-ico" aria-hidden="true">{icon}</span>{esc(label)}</a>'
        )
    return "".join(parts)


def display_html():
    return '''<div class="display-pop" hidden>
<p id="display-title">Display</p>
<div class="font-row" role="group" aria-label="Font size">
<button type="button" data-font-set="sm" aria-label="Smaller text">A−</button>
<button type="button" data-font-set="md" aria-label="Default text">A</button>
<button type="button" data-font-set="lg" aria-label="Larger text">A+</button>
<button type="button" data-font-set="xl" aria-label="Largest text">A++</button>
</div>
<label class="ink-label">Font colour <input data-ink type="color" value="#1e1b4b" aria-label="Font colour"></label>
<div class="theme-row" role="group" aria-label="Theme">
<button type="button" data-set-theme="bright">Bright</button>
<button type="button" data-set-theme="dark">Dark</button>
</div>
<label class="ink-label"><input data-reduce-motion type="checkbox"> Reduce motion</label>
</div>'''


BOOT = """<script>try{var s=JSON.parse(localStorage.getItem("acl-display")||"{}");document.documentElement.dataset.theme=s.theme==="dark"?"dark":"bright";if(s.font)document.documentElement.dataset.font=s.font;if(s.ink&&/^#[0-9a-fA-F]{6}$/.test(s.ink))document.documentElement.style.setProperty("--ink",s.ink);if(s.motion==="reduce")document.documentElement.dataset.motion="reduce";}catch(e){document.documentElement.dataset.theme="bright";}</script>"""


def page(path: Path, title, description, depth, current, body, scripts=""):
    prefix = "../" * depth
    home_tour = ' data-tour="home"' if current == "home" else ""
    doc = f'''<!DOCTYPE html>
<html lang="en" data-theme="bright">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>{esc(title)}</title>
<meta name="description" content="{esc(description)}">
<meta name="theme-color" content="#4c1d95">
{BOOT}
<link rel="stylesheet" href="{prefix}assets/css/wow.css">
<link rel="stylesheet" href="{prefix}assets/css/frame.css">
<link rel="stylesheet" href="{prefix}assets/css/footer.css">
<link rel="stylesheet" href="{prefix}guide/guide.css">
<link rel="icon" href="{prefix}favicon.svg" type="image/svg+xml">
<link rel="manifest" href="{prefix}site.webmanifest">
</head>
<body data-page="{esc(current)}">
<a class="skip" href="#main">Skip to content</a>
<header class="topbar">
<button type="button" class="menu-btn" data-open-menu data-tour="menu" aria-controls="site-drawer" aria-expanded="false">☰ Menu</button>
<a class="brand-home" id="tour-home" href="{prefix}index.html"{home_tour}>Connect Lab<small>Amazon Connect study</small></a>
<div class="rail-wrap">
<button type="button" class="rail-arrow" data-rail="-1" aria-label="Scroll sections left">‹</button>
<nav class="section-rail" aria-label="Sections">{banner_html(depth, current)}</nav>
<button type="button" class="rail-arrow" data-rail="1" aria-label="Scroll sections right">›</button>
</div>
<button type="button" class="display-btn" data-open-display aria-expanded="false" aria-controls="display-title">Display</button>
</header>
<div class="drawer-scrim" data-close-drawer></div>
<aside class="drawer" id="site-drawer" aria-hidden="true" inert aria-label="Menu">{drawer_html(depth, current)}</aside>
{display_html()}
<main id="main">{body}</main>
{footer_block()}
<nav class="phone-dock" aria-label="Quick">{dock_html(depth)}</nav>
<div id="guide-mount"></div>
<script src="{prefix}assets/js/frame.js" defer></script>
<script src="{prefix}assets/js/university.js" defer></script>
{scripts}
<script type="module" src="{prefix}guide/guide.js"></script>
</body>
</html>
'''
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(doc, encoding="utf-8")


STANDALONE_STYLE = """<style>
body{margin:0;font-family:system-ui,sans-serif;background:#fff7ed;color:#1e1b4b}
.standalone{max-width:900px;margin:0 auto;padding:20px 16px 32px}
.disclaimer-panel{display:flex;gap:12px;text-align:left;padding:16px;border:2px solid #0f766e;border-radius:16px;background:linear-gradient(135deg,#dcfce7,#dbeafe);color:#0b1f2a}
.disclaimer-panel h2{margin:0 0 6px;font-size:1.1rem}
a{color:#0b3a8f}
svg.fig{width:100%;height:auto}
.brand-row{display:flex;gap:12px;justify-content:center;margin:12px auto;padding:10px;width:fit-content;border-radius:18px;border:2px solid #7c3aed;box-shadow:0 0 16px rgba(124,58,237,.35)}
.brand-icon{width:56px;height:56px;display:grid;place-items:center;border-radius:16px;color:#fff;text-decoration:none}
.brand-icon svg{width:36px;height:36px}
.brand-web{background:radial-gradient(circle at 30% 25%,#c4b5fd,#312e81 72%)}
.brand-gh{background:#24292f}
.brand-tree{background:radial-gradient(circle at 32% 24%,#d9f99d,#4ade80 42%,#166534 78%);box-shadow:0 0 0 3px #052e16,0 0 18px #4ade80}
.disclaimer-panel a{color:#1e3a8a;font-weight:800}
</style>"""


def standalone(path: Path, title, body, depth):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        f'<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width, initial-scale=1"><title>{esc(title)}</title>{STANDALONE_STYLE}</head>'
        f'<body><article class="standalone"><div class="disclaimer-panel" role="note"><div><h2>Independent study lab</h2><p>{esc(DISCLAIMER)}</p>'
        f'<p>Contact: <a href="{CONTACT}">{CONTACT}</a></p></div></div><h1>{esc(title)}</h1>{body}{footer_block()}</article></body></html>',
        encoding="utf-8",
    )


def dl(depth, kind, href, label):
    return f'<a class="dl-btn dl-{kind}" href="{esc(rel(depth, href))}" download>{esc(label)}</a>'


def section_file(slug, section, title, inner):
    path = ROOT / "downloads" / "sections" / f"{slug}-{section}.html"
    standalone(path, title, inner, 2)
    return f"downloads/sections/{slug}-{section}.html"


def block(depth, sid, heading, inner, download):
    button = dl(depth, "html", download, "HTML")
    return f'<section class="block" id="{sid}" data-read><div class="section-bar"><h2>{esc(heading)}</h2>{button}</div>{inner}</section>'


def write_asset(rel_path, text):
    path = ROOT / rel_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8")
    return rel_path


def rsvg(svg_path, png_path):
    png_path.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run(["rsvg-convert", "-w", "1400", "-o", str(png_path), str(svg_path)], check=True)


def chrome(args, expect_file=None, timeout=25):
    """Headless Chrome writes the file, then keeps a background client alive. Kill it once the file is ready."""
    folder = tempfile.mkdtemp(prefix="cchrome-")
    cmd = [
        "google-chrome", "--headless=new", "--disable-gpu", "--no-sandbox", "--disable-dev-shm-usage",
        "--no-first-run", "--no-default-browser-check", "--disable-background-networking",
        "--disable-sync", "--disable-extensions", "--disable-component-update",
        f"--user-data-dir={folder}",
    ] + args
    proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=subprocess.DEVNULL, start_new_session=True)
    chunks = []

    def read_stdout():
        chunks.append(proc.stdout.read())

    reader = threading.Thread(target=read_stdout)
    reader.start()
    ready = False
    deadline = time.time() + timeout
    try:
        while time.time() < deadline:
            if expect_file:
                path = Path(expect_file)
                if path.exists() and path.stat().st_size > 400:
                    time.sleep(0.25)
                    ready = path.exists() and path.stat().st_size > 400
                    if ready:
                        break
            if proc.poll() is not None:
                ready = True
                break
            time.sleep(0.15)
        if expect_file and not ready:
            raise SystemExit("chrome did not write " + str(expect_file))
    finally:
        if proc.poll() is None:
            os.killpg(proc.pid, signal.SIGKILL)
            proc.wait(timeout=5)
        reader.join(timeout=2)
        shutil.rmtree(folder, ignore_errors=True)
    stdout = b"".join(chunks).decode("utf-8", "replace")
    return subprocess.CompletedProcess(cmd, 0, stdout, "")


def chrome_pdf(html_path, pdf_path):
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    if pdf_path.exists():
        pdf_path.unlink()
    chrome(["--no-pdf-header-footer", f"--print-to-pdf={pdf_path.resolve()}", html_path.resolve().as_uri()], expect_file=pdf_path)


def chrome_png(html_path, png_path):
    png_path.parent.mkdir(parents=True, exist_ok=True)
    if png_path.exists():
        png_path.unlink()
    chrome(["--window-size=1200,1600", "--hide-scrollbars", f"--screenshot={png_path.resolve()}", html_path.resolve().as_uri()], expect_file=png_path)
    fallback = Path.cwd() / "screenshot.png"
    if not png_path.exists() and fallback.exists():
        fallback.replace(png_path)


def prose(paragraphs):
    return "".join(f"<p>{esc(p)}</p>" for p in paragraphs)


def bullets(items):
    return "<ul>" + "".join(f"<li>{esc(item)}</li>" for item in items) + "</ul>"


def quiz_html(questions):
    parts = []
    for index, item in enumerate(questions, 1):
        choices = "".join(f"<button type=\"button\">{esc(choice)}</button>" for choice in item["choices"])
        parts.append(
            f'<article data-quiz data-answer="{int(item["answer"])}"><h3>Question {index}</h3><p>{esc(item["q"])}</p>'
            f'<div class="quiz-options">{choices}</div><details><summary>Show the cited answer</summary>'
            f'<p>{esc(item["why"])}</p><p><a href="{esc(item["url"])}">{esc(item["source"])}</a></p></details></article>'
        )
    return "".join(parts)


def build():
    grouped = by_track()
    known = lesson_map()
    quiz_pool = defaultdict(list)
    for item in load_quiz():
        quiz_pool[item["track"]].append(item)
    terms = defaultdict(list)
    for row in GLOSSARY:
        terms[row[2]].append(row)

    if (ROOT / "rooms").exists():
        shutil.rmtree(ROOT / "rooms")
    if (ROOT / "tracks").exists():
        shutil.rmtree(ROOT / "tracks")
    if (ROOT / "downloads").exists():
        shutil.rmtree(ROOT / "downloads")
    for folder in (ROOT / "assets" / "diagrams", ROOT / "assets" / "flows"):
        if folder.exists():
            shutil.rmtree(folder)

    diagram_files = []
    flow_files = []

    def add_diagram(slug, title, svg, animated=False):
        svg_rel = f"assets/diagrams/{slug}.svg"
        png_rel = f"assets/diagrams/{slug}.png"
        write_asset(svg_rel, svg)
        html_rel = f"downloads/diagrams/{slug}.html"
        standalone(ROOT / html_rel, title, f'<img class="fig" alt="" src="../../{svg_rel}" width="1400" height="800">', 2)
        diagram_files.append({"slug": slug, "title": title, "svg": svg_rel, "png": png_rel, "html": html_rel, "animated": animated})
        return svg_rel, png_rel, html_rel

    def add_flow(slug, title, svg):
        svg_rel = f"assets/flows/{slug}.svg"
        png_rel = f"assets/flows/{slug}.png"
        write_asset(svg_rel, svg)
        html_rel = f"downloads/flows/{slug}.html"
        standalone(ROOT / html_rel, title, f'<img class="fig" alt="" src="../../{svg_rel}" width="1400" height="800">', 2)
        flow_files.append({"slug": slug, "title": title, "svg": svg_rel, "png": png_rel, "html": html_rel})
        return svg_rel, png_rel, html_rel

    hero_steps = [
        {"label": "PSTN", "detail": "The caller dials a claimed number."},
        {"label": "Contact flow", "detail": "The inbound flow runs."},
        {"label": "Lex or queue", "detail": "Get customer input, or set the working queue."},
        {"label": "Routing profile", "detail": "Priority and delay decide eligibility."},
        {"label": "CCP", "detail": "The browser softphone uses TCP 443 and UDP 3478."},
        {"label": "CTR and Kinesis", "detail": "The finished contact can leave on a stream."},
    ]
    add_flow("inbound-call-path", "Inbound call path", flow_svg("Inbound call path", hero_steps))
    add_diagram("inbound-call-path", "Inbound call path", path_svg("Inbound call path", [{"label": s["label"]} for s in hero_steps], "Study picture. Confirm the cited rooms.", True), True)

    for track in TRACKS:
        slug = f"mind-{track['id']}"
        add_diagram(slug, track["title"] + " mind map", mind_svg(track, grouped[track["id"]]))

    room_meta = []
    for lesson in LESSONS:
        svg = diagram_svg(lesson["title"], lesson["diagram"], animated=False)
        flow = flow_svg(lesson["title"] + " flow", lesson["flow"])
        d_svg, d_png, d_html = add_diagram(lesson["id"], lesson["title"], svg)
        f_svg, f_png, f_html = add_flow(lesson["id"], lesson["title"] + " flow", flow)
        questions = pick_quiz(lesson, quiz_pool[lesson["track"]])
        track = next(item for item in TRACKS if item["id"] == lesson["track"])
        confirm = '<p class="confirm">Confirm before relying. Prices, quotas, and exam names change. Use the source linked on this page.</p>' if confirm_needed(lesson) else ""
        learn_inner = f'<p class="kicker">About five minutes of reading</p><p>{esc(lesson["why"])}</p>{prose(lesson["paras"])}{bullets(lesson["points"])}<p><strong>Number on this page:</strong> {esc(lesson["number"]["k"])} {esc(lesson["number"]["v"])}. {esc(lesson["number"]["n"])}</p>{confirm}'
        visual_inner = (
            f'<div class="visual"><img src="{esc(rel(1, d_svg))}" width="960" height="280" alt="{esc(lesson["title"])}"></div>'
            f'<div class="dl-row">{dl(1, "svg", d_svg, "Diagram SVG")}{dl(1, "png", d_png, "Diagram PNG")}{dl(1, "html", d_html, "Diagram HTML")}'
            f'{dl(1, "svg", f_svg, "Flow SVG")}{dl(1, "png", f_png, "Flow PNG")}{dl(1, "html", f_html, "Flow HTML")}</div>'
            f'<p>{esc(lesson["diagram"].get("caption") or "Picture from this room.")}</p>'
        )
        example_inner = "<ol>" + "".join(f"<li><strong>{esc(step['label'])}.</strong> {esc(step['detail'])}</li>" for step in lesson["flow"]) + "</ol>"
        example_inner += f'<p><strong>What usually goes wrong:</strong> {esc(lesson["mistake"])}</p>'
        source_inner = f'<p><a href="{esc(lesson["cite"]["url"])}">{esc(lesson["cite"]["title"])}</a></p><p>The prose is a study paraphrase. The link is the source of the facts.</p>'
        glossary_bits = terms[lesson["track"]][:4]
        gloss = "<ul>" + "".join(f'<li><strong>{esc(row[0])}</strong> {esc(row[1])}</li>' for row in glossary_bits) + "</ul>" if glossary_bits else "<p>See the glossary for this track.</p>"
        iac = IAC.get(lesson["track"])
        config = ""
        if iac:
            links = " ".join(f'<a href="{esc(url)}">{esc(name)}</a>' for name, url in iac["urls"])
            config = f'<p>{esc(iac["blurb"])}</p><p>{links}</p><details><summary>CloudFormation study snippet</summary><pre>{esc(iac["cfn"])}</pre></details>'
        future = known["eco-future"]["paras"][0]
        eco_inner = (
            f"<h3>Glossary</h3>{gloss}<h3>Mind map</h3><p><a href=\"{esc(rel(1, 'assets/diagrams/mind-' + lesson['track'] + '.svg'))}\">Open the {esc(track['title'])} mind map</a></p>"
            f"<h3>Sister technologies</h3><p>{esc(track['why'])}</p><h3>Config and troubleshooting</h3><p>{esc(lesson['mistake'])}</p>{config}"
            f"<h3>Integration map</h3><p>Join records on ContactId. The integration map room lists the seams.</p>"
            f"<h3>Cross-platform bridges</h3><p>{esc(known['br-map']['paras'][0])}</p><p><a href=\"{esc(rel(1, 'rooms/br-map.html'))}\">Open the mapping room</a></p>"
            f"<h3>Future hooks</h3><p>{esc(future)}</p>"
        )
        sections = {
            "learn": ("Learn it in 5 minutes", learn_inner),
            "visual": ("Visual", visual_inner),
            "example": ("Worked example", example_inner),
            "quiz": ("Quiz", quiz_html(questions)),
            "sources": ("Sources", source_inner),
            "ecosystem": ("Ecosystem", eco_inner),
        }
        blocks = []
        for sid, (heading, inner) in sections.items():
            href = section_file(lesson["id"], sid, lesson["title"] + " — " + heading, inner)
            blocks.append(block(1, sid, heading, inner, href))
        body = f'<div class="wrap"><p class="kicker">{esc(track["title"])}</p><h1>{esc(lesson["title"])}</h1>{"".join(blocks)}</div>'
        page(ROOT / "rooms" / f"{lesson['id']}.html", lesson["title"] + " · Connect Lab", lesson["why"], 1, "room", body)
        room_meta.append({"id": lesson["id"], "title": lesson["title"], "why": lesson["why"], "track": lesson["track"], "href": f"rooms/{lesson['id']}.html"})

    for track in TRACKS:
        cards = []
        for lesson in grouped[track["id"]]:
            cards.append(f'<a class="card" data-card href="{esc(rel(2, "rooms/" + lesson["id"] + ".html"))}"><h2>{esc(lesson["title"])}</h2><p>{esc(lesson["why"])}</p></a>')
        learn = f'<p>{esc(track["why"])}</p><p>{len(grouped[track["id"]])} rooms. Open one, then use its quiz.</p>'
        href = section_file("track-" + track["id"], "learn", track["title"], learn)
        body = f'<div class="wrap"><p class="kicker">Track</p><h1>{esc(track["title"])}</h1>{block(2, "learn", "Learn it in 5 minutes", learn, href)}<div class="card-grid">{"".join(cards)}</div></div>'
        page(ROOT / "tracks" / track["id"] / "index.html", track["title"] + " · Connect Lab", track["why"], 2, "learn", body)

    # PNG for every svg via rsvg
    for item in diagram_files + flow_files:
        rsvg(ROOT / item["svg"], ROOT / item["png"])

    write_hubs(grouped, known, room_meta, diagram_files, flow_files, quiz_pool)
    print(f"rooms {len(room_meta)} diagrams {len(diagram_files)} flows {len(flow_files)}")


def write_hubs(grouped, known, room_meta, diagram_files, flow_files, quiz_pool):
    track_cards = []
    for index, track in enumerate(TRACKS, 1):
        track_cards.append(
            f'<a class="track-card" data-card style="--hue:{esc(track["hue"])}" href="tracks/{esc(track["id"])}/index.html">'
            f'<p class="kicker">{index:02d}</p><h3>{esc(track["title"])}</h3><p>{esc(track["kicker"])}</p><p>{len(grouped[track["id"]])} rooms</p></a>'
        )
    hero = (ROOT / "assets" / "diagrams" / "inbound-call-path.svg").read_text(encoding="utf-8")
    home_learn = "<p>Amazon Connect here means the contact-center instance: flows, queues, routing profiles, Lex, agent assist, Contact Lens, Lambda, Kinesis, the CCP, and the bridges from UCCE, Genesys, and Webex Contact Center.</p><p>No AWS account is required. Nothing on this site places a live call.</p>"
    home_sections = [
        block(0, "learn", "Learn it in 5 minutes", home_learn, section_file("home", "learn", "Home", home_learn)),
    ]
    stats = f'''<section class="stats" aria-label="Counts in this repository">
<div class="stat"><b>{len(room_meta)}</b><span>rooms</span></div>
<div class="stat"><b>{len(TRACKS)}</b><span>tracks</span></div>
<div class="stat"><b>{len(diagram_files)}</b><span>diagrams</span></div>
<div class="stat"><b>{len(flow_files)}</b><span>flows</span></div>
</section>'''
    home = f'''<div class="wrap"><section class="hero" data-read>
<div><p class="kicker">Amazon Connect study · offline</p><h1>Learn the contact center in color.</h1>
<p class="lede">Flows, queues, routing profiles, Lex, agent assist, Contact Lens, Lambda, Kinesis, the CCP, and the bridges from Cisco UCCE, Genesys, and Webex.</p>
<div class="actions"><a class="btn" href="start.html">Start here</a><a class="ghost" href="#rooms">All rooms</a><a class="ghost" href="downloads.html">Downloads</a><a class="ghost" href="play.html">Play</a></div>
</div><div class="hero-visual">{hero}</div></section>
{"".join(home_sections)}{stats}
<section class="block" id="rooms" data-read><div class="section-bar"><h2>Tracks</h2>{dl(0, "html", section_file("home", "tracks", "Tracks", "<p>Eighteen tracks.</p>"), "HTML")}</div>
<label class="menu-search">Search rooms <input data-filter-cards="#rooms" type="search" aria-label="Search rooms"></label>
<div class="track-grid">{"".join(track_cards)}</div></section></div>'''
    page(ROOT / "index.html", "Amazon Connect Lab", "Independent study of Amazon Connect flows, queues, routing, and bridges.", 0, "home", home)

    ten = ["foundations-why", "flows-types", "route-queue", "route-profile", "foundations-lambda"]
    ten_html = "<ol>" + "".join(f'<li><a href="rooms/{item}.html">{esc(known[item]["title"])}</a> — {esc(known[item]["why"])}</li>' for item in ten) + "</ol>"
    start_body = f'''<div class="wrap"><h1>Start here</h1>
{block(0, "ten", "Finish in 10 minutes", ten_html + "<p>Five rooms. Then stop or keep going.</p>", section_file("start", "ten", "Ten minutes", ten_html))}
{block(0, "full", "Full course", "<p>Open a track from Home. The order on the mastery page runs from foundations to the ecosystem.</p><p><a href=\"index.html#rooms\">All rooms</a></p>", section_file("start", "full", "Full course", "<p>Use the track map.</p>"))}
{block(0, "download", "Just download", "<p><a href=\"downloads.html\">Open downloads</a> for sheets, cards, templates, the deck, and the zip.</p>", section_file("start", "download", "Just download", "<p>Downloads.</p>"))}
</div>'''
    page(ROOT / "start.html", "Start here · Connect Lab", "A ten-minute path and the full course.", 0, "learn", start_body)

    def gallery(files, kind):
        cards = []
        for item in files:
            buttons = dl(0, "svg", item["svg"], "SVG") + dl(0, "png", item["png"], "PNG") + dl(0, "html", item["html"], "HTML")
            cards.append(f'<article class="card" data-card><h3>{esc(item["title"])}</h3><img src="{esc(item["svg"])}" width="640" height="360" alt="{esc(item["title"])}"><div class="dl-row">{buttons}</div></article>')
        return "".join(cards)

    diagrams_body = f'<div class="wrap"><h1>Diagrams</h1><p>{len(diagram_files)} pictures. Each one is SVG, PNG, and a standalone HTML page.</p><div class="card-grid">{gallery(diagram_files, "diagram")}</div></div>'
    page(ROOT / "diagrams.html", "Diagrams · Connect Lab", "Diagrams from the Amazon Connect rooms.", 0, "diagrams", diagrams_body)
    flows_body = f'<div class="wrap"><h1>Flows</h1><p>{len(flow_files)} animated flows. The moving dot is a study aid.</p><div class="card-grid">{gallery(flow_files, "flow")}</div></div>'
    page(ROOT / "flows.html", "Flows · Connect Lab", "Animated flows from the Amazon Connect rooms.", 0, "flows", flows_body)

    def bar_chart(title, pairs):
        peak = max(value for _, value in pairs) or 1
        rows = []
        width, height = 760, 36 + len(pairs) * 36
        for i, (label, value) in enumerate(pairs):
            y = 16 + i * 36
            bar = 40 + (value / peak) * 420
            rows.append(f'<text x="8" y="{y + 18}" font-size="13" font-family="system-ui,sans-serif" fill="#1e1b4b">{esc(label)}</text>')
            rows.append(f'<rect x="180" y="{y}" width="{bar:.1f}" height="22" rx="6" fill="{FILLS[i % 6]}" stroke="{STROKES[i % 6]}"/>')
            rows.append(f'<text x="{184 + bar:.1f}" y="{y + 16}" font-size="13" font-family="system-ui,sans-serif" fill="#1e1b4b">{value}</text>')
        return f'<section class="block" data-chart="{esc(title)}"><h2>{esc(title)}</h2><svg viewBox="0 0 {width} {height}" width="760" height="{height}" role="img" aria-label="{esc(title)}">{"".join(rows)}</svg><p>These are counts of files and lessons in this repository. They are not traffic, revenue, or a quota.</p></section>'

    lesson_pairs = [(track["title"], len(grouped[track["id"]])) for track in TRACKS]
    point_pairs = [(track["title"], sum(len(item["points"]) for item in grouped[track["id"]])) for track in TRACKS]
    flow_pairs = [(track["title"], sum(len(item["flow"]) for item in grouped[track["id"]])) for track in TRACKS]
    term_pairs = [(track["title"], sum(1 for row in GLOSSARY if row[2] == track["id"])) for track in TRACKS]
    quiz_pairs = [(track["title"], len(quiz_pool[track["id"]])) for track in TRACKS]
    kind_pairs = []
    kinds = defaultdict(int)
    for lesson in LESSONS:
        kinds[lesson["diagram"]["type"]] += 1
    kind_pairs = list(kinds.items())
    charts = "".join([
        bar_chart("Lessons in each track", lesson_pairs),
        bar_chart("Teaching points in each track", point_pairs),
        bar_chart("Flow steps in each track", flow_pairs),
        bar_chart("Glossary terms in each track", term_pairs),
        bar_chart("Practice questions in each track", quiz_pairs),
        bar_chart("Diagram shapes", kind_pairs),
    ])
    page(ROOT / "charts.html", "Charts · Connect Lab", "Repository counts. Not traffic or revenue.", 0, "charts", f'<div class="wrap"><h1>Charts</h1>{charts}</div>')

    games = f'''<div class="wrap"><h1>Play</h1>
<p>These drills use sentences from the rooms. They do not open Amazon Connect.</p>
{game("Priority", "A lower priority number is offered first. Delay 0 is eligible immediately. Source: the routing profile room.", [("Orders at priority 1, delay 0", 0), ("Overflow at priority 2, delay 0", 1)], 0, "rooms/route-profile.html")}
{game("Lambda", "Synchronous invoke waits at most 8 seconds. A 30 second Lambda console timeout does not make the flow wait.", [("The flow can take the timeout branch", 0), ("The flow waits the full 30 seconds", 1)], 0, "rooms/foundations-lambda.html")}
{game("Bridge", "A UCCE skill group maps to a queue. A CUCM partition is not a Connect object.", [("Queue", 0), ("A CUCM partition object inside Connect", 1)], 0, "rooms/br-map.html")}
{game("Ports", "The network lesson names the instance on TCP 443 and TurnNlb on UDP 3478. Opening only 443 is the mistake in that room.", [("UDP 3478 to TurnNlb", 0), ("TCP 443 is enough for softphone media", 1)], 0, "rooms/foundations-network.html")}
</div>'''
    page(ROOT / "play.html", "Play · Connect Lab", "Short drills cited to the rooms.", 0, "play", games, '<script src="assets/js/play.js" defer></script>')

    listen = f'''<div class="wrap"><h1>Listen</h1>
{block(0, "clips", "Study clips", "<p>Two short synthesized clips. They are not a recording of a person. The script is in the audio folder.</p><audio controls preload=\"none\" src=\"audio/use.mp3\"></audio><p><a href=\"audio/scripts/architect-use.md\">Script for the first clip</a></p><audio controls preload=\"none\" src=\"audio/share.mp3\"></audio><p><a href=\"audio/scripts/architect-share.md\">Script for the second clip</a></p>", section_file("listen", "clips", "Listen", "<p>Synthesized study clips.</p>"))}
{block(0, "read", "Read-along", "<p>Open the monkey, choose Read this page or Read along, and pick a voice. The current section glows and the spoken sentence is marked in the bot. The mic uses the browser speech recognizer, with the text box when the mic is missing.</p>", section_file("listen", "read", "Read-along", "<p>Browser voice.</p>"))}
</div>'''
    page(ROOT / "listen.html", "Listen · Connect Lab", "Study clips and read-along.", 0, "listen", listen)

    help_body = '''<div class="wrap"><h1>Help</h1>
<section class="block" id="menu" data-read><h2>Menu</h2><p>The Menu button is at the top left. The drawer slides in. Each row has an icon, a label, and one line. Search filters those rows. Escape or the backdrop closes it.</p><p>On a phone the bottom bar is Menu, Learn, Diagrams, Play, Listen, and Guide.</p></section>
<section class="block" id="course" data-read><h2>Course</h2><p>Start with the ten-minute path, or open the full track map. Every room has a five-minute read, a picture, a worked example, a quiz, sources, and an ecosystem panel.</p></section>
<section class="block" id="bot" data-read><h2>Bot</h2><p>The monkey greets you once after the tour, then docks. Drag it. The position is remembered in this browser. After a few quiet minutes it may offer help, up to three times. Not now waits. Stop offering remembers that choice.</p><p>Finish in 10 minutes, Full course, and Just download are the first choices. You can also ask it to open a room.</p></section>
<section class="block" id="downloads" data-read><h2>Downloads</h2><p>Colored buttons sit in every section. The downloads page collects sheets, cards, templates, the deck, and one zip.</p></section>
<section class="block" id="tour" data-read><h2>Tour</h2><p>The first visit walks six steps: Home, Menu, Learn, Diagrams, Play, and Listen. Downloads and the bot are highlighted along the way. Next, Back, and Skip are on the card.</p><p><button type="button" data-replay-tour>Replay the tour</button></p></section>
<section class="block" id="display" data-read><h2>Display</h2><p>Display, on the right of the banner, changes font size, font colour, bright or dark, and motion.</p></section>
</div>'''
    page(ROOT / "help.html", "Help · Connect Lab", "How to use the menu, the course, the bot, and the downloads.", 0, "guide", help_body)

    legal = f'''<div class="wrap"><h1>Legal and credits</h1>
<section class="block" data-read><h2>Copyright</h2><p>The code in this repository is under the MIT license in the LICENSE file. NOTICE.md describes the study text as CC BY-NC-SA 4.0. README.md has also called that text MIT. Treat NOTICE.md as the statement for the study text until those two files are reconciled. Confirm before relying on either sentence for a reuse decision.</p></section>
<section class="block" data-read><h2>Trademarks</h2><p>Amazon Web Services, Amazon Connect, Amazon Lex, Amazon CloudWatch, and other AWS names are trademarks of their owners. Cisco, Webex, Avaya, Genesys, Salesforce, ServiceNow, and Microsoft names belong to their owners. This lab is not affiliated with or endorsed by them.</p></section>
<section class="block" data-read><h2>Third-party licences</h2><p>The published pages do not load a third-party runtime. The older builder under web/vendor still contains Three.js (MIT), anime.js, D3 (BSD), Rough.js, and canvas-confetti. Their license files stay next to those files. This page does not redistribute their source.</p></section>
<section class="block" data-read><h2>Privacy</h2><p>There is no account and no tracker. Theme, tour, quiz, and bot position stay in this browser's localStorage. The service worker caches pages you open on this origin.</p></section>
<section class="block" data-read><h2>Accessibility</h2><p>The banner has a skip link. Display changes font size and font colour. The drawer and the bot work from the keyboard. Reduced motion is in Display and follows the system setting.</p></section>
<section class="block" data-read><h2>Removal</h2><p>A concern can be sent through <a href="{CONTACT}">{CONTACT}</a>. The page will be removed or corrected. Nothing here is sold.</p></section>
</div>'''
    page(ROOT / "legal.html", "Legal · Connect Lab", "Copyright, trademarks, privacy, and removal.", 0, "legal", legal)

    eco = f'''<div class="wrap"><h1>Ecosystem</h1>
{block(0, "glossary", "Glossary", f'<p><a href="glossary.html">{len(GLOSSARY)} terms</a> collected from the rooms.</p>', section_file("ecosystem", "glossary", "Glossary", "<p>See the glossary.</p>"))}
{block(0, "mind", "Mind map", '<p><a href="assets/diagrams/mind-ecosystem.svg">Ecosystem mind map</a></p>', section_file("ecosystem", "mind", "Mind map", "<p>Mind map.</p>"))}
{block(0, "sisters", "Sister technologies", f'<p>{esc(known["eco-why"]["why"])}</p>{prose(known["eco-why"]["paras"])}', section_file("ecosystem", "sisters", "Sisters", prose(known["eco-why"]["paras"])))}
{block(0, "config", "Config and troubleshooting", f'<p>{esc(known["ts-method"]["why"])}</p>{bullets(known["ts-method"]["points"])}', section_file("ecosystem", "config", "Troubleshooting", bullets(known["ts-method"]["points"])))}
{block(0, "map", "Integration map", prose(known["eco-map"]["paras"]) + bullets(known["eco-map"]["points"]), section_file("ecosystem", "map", "Integration map", prose(known["eco-map"]["paras"])))}
{block(0, "bridges", "Cross-platform bridges", prose(known["br-map"]["paras"]) + bullets(known["br-map"]["points"]), section_file("ecosystem", "bridges", "Bridges", prose(known["br-map"]["paras"])))}
{block(0, "future", "Future hooks", prose(known["eco-future"]["paras"]), section_file("ecosystem", "future", "Future hooks", prose(known["eco-future"]["paras"])))}
</div>'''
    page(ROOT / "ecosystem.html", "Ecosystem · Connect Lab", "Glossary, mind map, bridges, and hooks.", 0, "ecosystem", eco)

    gloss_rows = "".join(
        f"<tr><td>{esc(row[0])}</td><td>{esc(row[1])}</td><td><a href=\"{esc(row[3])}\">{esc(row[4])}</a></td></tr>" for row in GLOSSARY
    )
    page(ROOT / "glossary.html", "Glossary · Connect Lab", "Terms from the written rooms.", 0, "glossary",
         f'<div class="wrap"><h1>Glossary</h1><div class="table-wrap"><table><thead><tr><th>Term</th><th>Meaning</th><th>Source</th></tr></thead><tbody>{gloss_rows}</tbody></table></div></div>')

    cards = [{"front": row[0], "back": row[1], "source": row[4], "url": row[3], "track": row[2]} for row in GLOSSARY]
    (ROOT / "assets" / "data").mkdir(parents=True, exist_ok=True)
    (ROOT / "assets" / "data" / "flashcards.json").write_text(json.dumps({"cards": cards}, indent=2), encoding="utf-8")
    (ROOT / "downloads").mkdir(parents=True, exist_ok=True)
    shutil.copyfile(ROOT / "assets" / "data" / "flashcards.json", ROOT / "downloads" / "flashcards.json")
    questions = []
    for track_items in quiz_pool.values():
        questions.extend(track_items)
    (ROOT / "assets" / "data" / "quiz.json").write_text(json.dumps({"questions": questions}, indent=2), encoding="utf-8")
    page(ROOT / "flashcards.html", "Flashcards · Connect Lab", "Know-or-again cards.", 0, "cards",
         f'<div class="wrap"><h1>Flashcards</h1><div class="dl-row">{dl(0, "json", "downloads/flashcards.json", "JSON")}{dl(0, "html", "downloads/flashcards.html", "HTML")}{dl(0, "pdf", "downloads/flashcards.pdf", "PDF")}</div><div data-deck="assets/data/flashcards.json" data-mode="cards"></div></div>',
         '<script src="assets/js/study.js" defer></script>')
    page(ROOT / "quizzes.html", "Quizzes · Connect Lab", "Practice questions with cited answers.", 0, "quizzes",
         '<div class="wrap"><h1>Quizzes</h1><p>Answers stay cited. A wrong choice is a study miss, not a score sent anywhere.</p><div data-deck="assets/data/quiz.json" data-mode="quiz"></div></div>',
         '<script src="assets/js/study.js" defer></script>')

    flash_html = "".join(f"<h2>{esc(card['front'])}</h2><p>{esc(card['back'])}</p><p>Source: {esc(card['source'])}</p>" for card in cards)
    standalone(ROOT / "downloads" / "flashcards.html", "Flashcards", flash_html, 1)

    write_sheets(known)
    write_templates(known)
    write_deck()
    write_reference_pages(known)
    write_downloads(diagram_files, flow_files)
    write_kb(room_meta)


def game(title, prompt, options, answer, href):
    radios = "".join(
        f'<label><input type="radio" name="pick" value="{index}"> {esc(label)}</label>' for label, index in options
    )
    return (
        f'<section class="block" data-read><h2>{esc(title)}</h2><p>{esc(prompt)}</p>'
        f'<form data-game data-answer="{answer}">{radios}<div class="actions"><button class="btn" type="submit">Check</button></div>'
        f'<p data-game-note></p><details><summary>Show the cited answer</summary><p><a href="{esc(href)}">Open the room</a></p></details></form></section>'
    )


def write_sheets(known):
    sheets = [
        ("ports-media", "Ports and media", ["foundations-network"]),
        ("flow-blocks", "Flow types and errors", ["flows-types", "flows-blocks", "flows-errors"]),
        ("routing-profile", "Queues and routing profiles", ["route-queue", "route-profile", "route-hours"]),
        ("lambda-contract", "Lambda contract", ["foundations-lambda", "lambda-contract"]),
        ("lex-input", "Lex in the flow", ["lex-get-input", "lex-slots", "lex-fallback"]),
        ("contact-lens", "Conversational analytics", ["an-lens", "eco-lens-ops"]),
        ("bridge-map", "UCCE, Genesys, and Webex bridges", ["br-map", "br-pitfalls"]),
        ("kinesis-ctr", "Kinesis and contact trace records", ["foundations-kinesis", "an-ctr"]),
        ("e911", "E911 for CCP agents", ["eco-e911"]),
        ("resiliency", "Global Resiliency pairs", ["foundations-regions", "res-pair"]),
        ("certs-badges", "Certifications and badges", ["foundations-certs"]),
        ("troubleshoot", "Troubleshooting habit", ["ts-method", "ts-faults"]),
    ]
    for slug, title, ids in sheets:
        parts = ['<p class="confirm">Confirm before relying. Quotas, prices, and names change.</p>']
        for item_id in ids:
            lesson = known[item_id]
            parts.append(f"<h2>{esc(lesson['title'])}</h2>{prose(lesson['paras'])}{bullets(lesson['points'])}<p>Source: {esc(lesson['cite']['title'])}</p>")
        standalone(ROOT / "downloads" / "sheets" / f"{slug}.html", title, "".join(parts), 2)


def write_templates(known):
    templates = [
        ("inbound-flow", "Inbound flow worksheet", ["flows-types", "flows-errors", "flows-attributes"]),
        ("routing-profile", "Routing profile worksheet", ["route-queue", "route-profile", "route-hours"]),
        ("cutover-checklist", "Cutover checklist", ["br-playbook", "br-map", "br-pitfalls"]),
        ("lambda-contract", "Lambda contract checklist", ["foundations-lambda", "lambda-contract", "lambda-perms"]),
    ]
    for slug, title, ids in templates:
        paragraphs = [DISCLAIMER, "Contact: " + CONTACT, title, "Confirm before relying."]
        html_parts = ['<p class="confirm">Confirm before relying.</p>']
        for item_id in ids:
            lesson = known[item_id]
            paragraphs.append(lesson["title"])
            paragraphs.extend(lesson["points"])
            paragraphs.append("Caution: " + lesson["mistake"])
            html_parts.append(f"<h2>{esc(lesson['title'])}</h2>{bullets(lesson['points'])}<p>{esc(lesson['mistake'])}</p>")
        standalone(ROOT / "downloads" / "templates" / f"{slug}.html", title, "".join(html_parts), 2)
        write_docx(ROOT / "downloads" / "templates" / f"{slug}.docx", title, paragraphs)


def write_docx(path: Path, title, paragraphs):
    def para(text, bold=False):
        weight = "<w:b/>" if bold else ""
        return f'<w:p><w:r><w:rPr>{weight}</w:rPr><w:t xml:space="preserve">{esc(text)}</w:t></w:r></w:p>'
    body = para(title, True) + "".join(para(text) for text in paragraphs)
    document = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
        '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body>'
        + body + "</w:body></w:document>"
    )
    types = '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Default Extension="xml" ContentType="application/xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>'''
    rels = '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>'''
    path.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", types)
        archive.writestr("_rels/.rels", rels)
        archive.writestr("word/document.xml", document)


def write_deck():
    prs = Presentation()
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = "Amazon Connect study"
    title_slide.placeholders[1].text = "Independent learning lab. " + DISCLAIMER
    slides = [
        ("Flows", "A contact flow is the program the caller hears. An unconnected error branch hangs up."),
        ("Queues and profiles", "A queue waits. A routing profile decides who is eligible. A lower priority number wins."),
        ("Lex", "Get customer input is where the bot is selected. Copy slots into attributes or the agent cannot see them."),
        ("Agent assist and Contact Lens", "Names moved in 2026. Confirm the current admin-guide name before you write a runbook."),
        ("Lambda", "Synchronous invoke waits at most 8 seconds. Return the shape the block expects."),
        ("Kinesis", "Agent events and contact trace records can leave on a stream. Join on ContactId."),
        ("CCP", "Sign-in can succeed when media cannot. The network lesson names TCP 443 and UDP 3478."),
        ("Bridges", "A UCCE skill group maps to a queue. A script maps to a contact flow. A partition is not a Connect object."),
        ("Confirm", "Prices, quotas, and exam names change. Confirm before relying. Nothing in this deck is an invoice."),
    ]
    for heading, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = heading
        frame = slide.placeholders[1].text_frame
        frame.text = body
        for paragraph in frame.paragraphs:
            paragraph.font.size = Pt(20)
    close = prs.slides.add_slide(prs.slide_layouts[1])
    close.shapes.title.text = "Independent study lab"
    close.placeholders[1].text = DISCLAIMER + " Contact: " + CONTACT
    path = ROOT / "downloads" / "connect-study-deck.pptx"
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)


def write_reference_pages(known):
    def room_page(name, current, title, lesson_ids):
        parts = [f'<div class="wrap"><h1>{esc(title)}</h1>']
        for item_id in lesson_ids:
            lesson = known[item_id]
            inner = prose(lesson["paras"]) + bullets(lesson["points"]) + f'<p>Source: <a href="{esc(lesson["cite"]["url"])}">{esc(lesson["cite"]["title"])}</a></p>'
            if confirm_needed(lesson):
                inner = '<p class="confirm">Confirm before relying.</p>' + inner
            href = section_file(name, item_id, lesson["title"], inner)
            parts.append(block(0, item_id, lesson["title"], inner, href))
        parts.append("</div>")
        page(ROOT / f"{name}.html", title + " · Connect Lab", title, 0, current, "".join(parts))

    room_page("certs", "certs", "Certifications", ["foundations-certs"])
    room_page("ports", "ports", "Ports and protocols", ["foundations-network"])
    room_page("regions", "regions", "Regions", ["foundations-regions", "res-pair"])
    mastery_items = []
    for track in TRACKS:
        mastery_items.append(f'<li><a href="tracks/{esc(track["id"])}/index.html">{esc(track["title"])}</a> — {esc(track["kicker"])}</li>')
    mastery = '<div class="wrap"><h1>Mastery</h1><p>This order is the lab sequence, from the AWS account under the instance through the bridges. It is not a certification path.</p><ol>' + "".join(mastery_items) + "</ol></div>"
    page(ROOT / "mastery.html", "Mastery · Connect Lab", "Fundamentals through the ecosystem.", 0, "mastery", mastery)
    interviews = '''<div class="wrap"><h1>Interview practice</h1>
<p class="confirm">These answers are fictional practice for a made-up team. They are not a biography and they are not a claim that anyone has run Amazon Connect in production.</p>
<section class="block" data-read><h2>Routing</h2><p>Situation: two teams were both set to priority 1. Task: keep overflow in reserve. Action: the fictional team set the primary queue to priority 1 with delay 0 and the reserve queue to a higher number. Result: the reserve team stopped taking the first call. Source: the routing profile room.</p></section>
<section class="block" data-read><h2>Lambda</h2><p>Situation: a function was tuned to 30 seconds in the Lambda console. Task: stop the flow from hanging up. Action: the fictional team matched the block timeout to work that can finish in 8 seconds, or moved the slow part to the async pattern in the lesson. Result: the error branch became a prompt instead of silence. Source: the Lambda room.</p></section>
<section class="block" data-read><h2>Bridge</h2><p>Situation: a skill group, a script, and a desktop gadget were all called the queue. Task: write a map. Action: the fictional team mapped the skill group to a queue, the script to a contact flow, and the desktop to the CCP. Result: the design review stopped arguing about one word. Source: the mapping room.</p></section>
</div>'''
    page(ROOT / "interviews.html", "Interviews · Connect Lab", "Fictional practice answers.", 0, "interviews", interviews)
    careers = '''<div class="wrap"><h1>Careers</h1>
<p class="confirm">These are study labels for this lab. They are not job requisitions, and they are not a biography.</p>
<ul><li>Flow reader: rooms on contact flows, modules, and error branches.</li>
<li>Routing reader: queues, hours, profiles, priority, and delay.</li>
<li>Telephony reader: number types, ports, and the CCP media path.</li>
<li>Bridge reader: the mapping room before a move from UCCE, Genesys, or Webex Contact Center.</li></ul>
<p>There is no Amazon Connect specialty exam in the certification room. Confirm that on the current AWS Certification catalog before you pay for anything.</p>
</div>'''
    page(ROOT / "careers.html", "Careers · Connect Lab", "Study paths. Not a job offer.", 0, "careers", careers)


def render_office_previews():
    for html_path in list((ROOT / "downloads" / "sheets").glob("*.html")) + list((ROOT / "downloads" / "templates").glob("*.html")):
        chrome_pdf(html_path, html_path.with_suffix(".pdf"))
        chrome_png(html_path, html_path.with_suffix(".png"))
    chrome_pdf(ROOT / "downloads" / "flashcards.html", ROOT / "downloads" / "flashcards.pdf")
    chrome_png(ROOT / "downloads" / "flashcards.html", ROOT / "downloads" / "flashcards.png")


def magic_ok(path: Path):
    data = path.read_bytes()[:8]
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return data.startswith(b"%PDF")
    if suffix == ".png":
        return data.startswith(b"\x89PNG")
    if suffix in {".docx", ".pptx", ".zip"}:
        return data.startswith(b"PK")
    if suffix == ".svg":
        return b"<svg" in path.read_bytes()[:200]
    if suffix == ".json":
        json.loads(path.read_text(encoding="utf-8"))
        return True
    if suffix in {".html", ".md"}:
        text = path.read_text(encoding="utf-8")
        return "learning-only" in text and "mailto:" not in text
    return path.stat().st_size > 0


def write_downloads(diagram_files, flow_files):
    render_office_previews()
    rows = []
    sample_html = [
        ROOT / "downloads" / "sheets" / "ports-media.html",
        ROOT / "downloads" / "diagrams" / "foundations-why.html",
        ROOT / "rooms" / "route-queue.html",
    ]
    dom_notes = []
    for sample in sample_html:
        result = chrome(["--dump-dom", sample.resolve().as_uri()], timeout=8)
        ok = "learning-only" in result.stdout
        dom_notes.append((sample.name, "pass" if ok else "fail"))
        if not ok:
            raise SystemExit("headless DOM missing disclaimer: " + sample.name)

    files = []
    for folder, kind in (
        (ROOT / "assets" / "diagrams", "svg"),
        (ROOT / "assets" / "flows", "svg"),
        (ROOT / "downloads", "mixed"),
    ):
        for path in folder.rglob("*"):
            if path.is_file() and path.suffix != ".zip":
                files.append(path)
    for path in files:
        ok = magic_ok(path)
        rows.append((str(path.relative_to(ROOT)), path.suffix.lstrip("."), "pass" if ok else "fail"))
        if not ok:
            raise SystemExit("bad file " + str(path))
    for name, status in dom_notes:
        rows.append((name + " headless DOM", "html", status))

    table = "<table><thead><tr><th>File</th><th>Kind</th><th>Check</th></tr></thead><tbody>" + "".join(
        f"<tr><td>{esc(name)}</td><td>{esc(kind)}</td><td>{esc(status)}</td></tr>" for name, kind, status in rows
    ) + "</tbody></table>"
    (ROOT / "docs").mkdir(exist_ok=True)
    (ROOT / "docs" / "download-tests.json").write_text(json.dumps([
        {"file": name, "kind": kind, "result": status} for name, kind, status in rows
    ], indent=2), encoding="utf-8")

    bundle = ROOT / "downloads" / "amazon-connect-lab-bundle.zip"
    with zipfile.ZipFile(bundle, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("README.txt", DISCLAIMER + "\nContact: " + CONTACT + "\n" + COPY + "\n")
        for path in (ROOT / "downloads").rglob("*"):
            if path.is_file() and path.suffix != ".zip":
                archive.write(path, path.relative_to(ROOT / "downloads"))
    rows.append((str(bundle.relative_to(ROOT)), "zip", "pass" if magic_ok(bundle) else "fail"))
    zip_size = bundle.stat().st_size

    open_buttons = (
        f'<p class="actions"><a class="btn" href="index.html">Open app</a> '
        f'<a class="btn" href="{SITE}/amazon-connect-lab-free-university-mesmerizing/">Open the published app</a> '
        f'<a class="ghost" href="{REPO}">View repo</a></p>'
    )
    groups = [
        ("Diagrams", diagram_files),
        ("Flows", flow_files),
    ]
    cards = []
    for label, items in groups:
        for item in items:
            cards.append(
                f'<article class="dl-card"><h3>{esc(item["title"])}</h3><p>{esc(label)}</p><div class="dl-row">'
                f'{dl(0, "svg", item["svg"], "SVG")}{dl(0, "png", item["png"], "PNG")}{dl(0, "html", item["html"], "HTML")}</div></article>'
            )
    sheet_cards = []
    for path in sorted((ROOT / "downloads" / "sheets").glob("*.html")):
        stem = path.stem
        sheet_cards.append(
            f'<article class="dl-card"><h3>{esc(stem.replace("-", " "))}</h3><div class="dl-row">'
            f'{dl(0, "html", "downloads/sheets/" + stem + ".html", "HTML")}'
            f'{dl(0, "pdf", "downloads/sheets/" + stem + ".pdf", "PDF")}'
            f'{dl(0, "png", "downloads/sheets/" + stem + ".png", "PNG")}</div></article>'
        )
    template_cards = []
    for path in sorted((ROOT / "downloads" / "templates").glob("*.html")):
        stem = path.stem
        template_cards.append(
            f'<article class="dl-card"><h3>{esc(stem.replace("-", " "))}</h3><div class="dl-row">'
            f'{dl(0, "html", "downloads/templates/" + stem + ".html", "HTML")}'
            f'{dl(0, "pdf", "downloads/templates/" + stem + ".pdf", "PDF")}'
            f'{dl(0, "docx", "downloads/templates/" + stem + ".docx", "DOCX")}</div></article>'
        )
    body = f'''<div class="wrap"><h1>Downloads</h1>
{open_buttons}
<p>Every file below is standalone and includes the disclaimer. The zip is {zip_size} bytes.</p>
<section class="block" id="bundle"><h2>Everything</h2><div class="dl-row">{dl(0, "zip", "downloads/amazon-connect-lab-bundle.zip", "ZIP")}</div></section>
<section class="block" id="sheets"><h2>Cheat sheets</h2><div class="dl-grid">{"".join(sheet_cards)}</div></section>
<section class="block" id="cards"><h2>Flashcards</h2><div class="dl-row">{dl(0, "json", "downloads/flashcards.json", "JSON")}{dl(0, "html", "downloads/flashcards.html", "HTML")}{dl(0, "pdf", "downloads/flashcards.pdf", "PDF")}{dl(0, "png", "downloads/flashcards.png", "PNG")}</div></section>
<section class="block" id="templates"><h2>Templates</h2><div class="dl-grid">{"".join(template_cards)}</div></section>
<section class="block" id="deck"><h2>Deck</h2><div class="dl-row">{dl(0, "pptx", "downloads/connect-study-deck.pptx", "PPTX")}</div></section>
<section class="block" id="pictures"><h2>Diagrams and flows</h2><div class="dl-grid">{"".join(cards)}</div></section>
<section class="block" id="tests"><h2>Headless test table</h2><p>Checked in this build: file types, the disclaimer string, and a headless DOM read of three pages.</p>{table}</section>
</div>'''
    page(ROOT / "downloads.html", "Downloads · Connect Lab", "Sheets, cards, templates, and the zip.", 0, "downloads", body)
    (ROOT / ".nojekyll").write_text("", encoding="utf-8")
    hosts = sorted({
        "cyphermonkeylabs.com", "github.com", "linktr.ee", "docs.aws.amazon.com", "aws.amazon.com",
        "hidekazu-konishi.com", "registry.terraform.io", "repost.aws", "www.servicenow.com",
    })
    (ROOT / "assets" / "data" / "allowed-hosts.json").write_text(json.dumps({"hosts": hosts}, indent=2), encoding="utf-8")


def write_kb(room_meta):
    pages = []
    for group, items in DRAWER:
        for key, label, href, icon, desc in items:
            pages.append({"title": label, "href": href.split("#")[0], "why": desc, "words": re.findall(r"[a-z]{4,}", (label + " " + desc).lower())})
    rooms = []
    for item in room_meta:
        words = re.findall(r"[a-z]{4,}", (item["title"] + " " + item["why"] + " " + item["track"]).lower())
        rooms.append({"title": item["title"], "href": item["href"], "why": item["why"], "words": words[:12]})
    (ROOT / "guide" / "kb.json").write_text(json.dumps({"pages": pages, "rooms": rooms}, indent=2), encoding="utf-8")


if __name__ == "__main__":
    build()
